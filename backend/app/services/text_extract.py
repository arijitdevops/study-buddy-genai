"""Text extraction from uploaded documents.

``pypdf`` and ``python-docx`` are synchronous and CPU-bound, so every call is
pushed onto a worker thread with :func:`asyncio.to_thread`. Blocking the event
loop here would stall every other request on the same worker.
"""

from __future__ import annotations

import asyncio
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

#: Paragraph-ish split used when chunking plain text.
_PARAGRAPH_RE = re.compile(r"\n\s*\n")
_WHITESPACE_RE = re.compile(r"[ \t]+")

DEFAULT_CHUNK_CHARS = 1400
DEFAULT_CHUNK_OVERLAP = 180


class ExtractionError(RuntimeError):
    """The file could not be read as text."""


@dataclass(slots=True)
class TextChunk:
    """One chunk of extracted text."""

    index: int
    content: str
    page: Optional[int] = None


@dataclass(slots=True)
class ExtractionResult:
    """The outcome of extracting text from one file."""

    text: str
    chunks: list[TextChunk] = field(default_factory=list)
    pages: int = 0

    @property
    def char_count(self) -> int:
        """Number of characters extracted."""
        return len(self.text)


def _normalise(text: str) -> str:
    """Collapse runs of spaces and trim trailing whitespace per line."""
    lines = [_WHITESPACE_RE.sub(" ", line).rstrip() for line in text.splitlines()]
    return "\n".join(lines).strip()


def _extract_pdf_sync(path: Path) -> tuple[str, list[tuple[int, str]]]:
    """Extract text from a PDF. Runs on a worker thread."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError(
            "pypdf is not installed. Run: pip install -r requirements.txt"
        ) from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001 - pypdf raises a wide range
        raise ExtractionError(f"This PDF could not be opened: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError("This PDF is password protected.") from exc

    pages: list[tuple[int, str]] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            content = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - a bad page should not fail the file
            logger.warning("Could not extract page %d; skipping.", number)
            continue
        content = _normalise(content)
        if content:
            pages.append((number, content))
    return "\n\n".join(text for _, text in pages), pages


def _extract_docx_sync(path: Path) -> str:
    """Extract text from a DOCX, including table cells. Runs on a thread."""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError(
            "python-docx is not installed. Run: pip install -r requirements.txt"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"This Word file could not be opened: {exc}") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _normalise("\n".join(parts))


def _extract_plain_sync(path: Path) -> str:
    """Read a text or markdown file, tolerating imperfect encodings."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return _normalise(path.read_text(encoding=encoding))
        except UnicodeDecodeError:
            continue
        except OSError as exc:
            raise ExtractionError(f"Could not read the file: {exc}") from exc
    raise ExtractionError("The file is not readable as text.")


def chunk_text(
    text: str,
    *,
    chunk_chars: int = DEFAULT_CHUNK_CHARS,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
    page: Optional[int] = None,
    start_index: int = 0,
) -> list[TextChunk]:
    """Split ``text`` into overlapping chunks on paragraph boundaries.

    Args:
        text: The text to split.
        chunk_chars: Target chunk size.
        overlap: Characters of tail carried into the next chunk.
        page: Page number to tag each chunk with.
        start_index: Index of the first chunk produced.

    Returns:
        A list of :class:`TextChunk`.
    """
    if not text.strip():
        return []

    paragraphs = [p.strip() for p in _PARAGRAPH_RE.split(text) if p.strip()]
    chunks: list[TextChunk] = []
    buffer = ""
    index = start_index

    def flush() -> None:
        nonlocal buffer, index
        if buffer.strip():
            chunks.append(TextChunk(index=index, content=buffer.strip(), page=page))
            index += 1

    for paragraph in paragraphs:
        if len(paragraph) > chunk_chars:
            flush()
            buffer = ""
            for start in range(0, len(paragraph), chunk_chars - overlap):
                piece = paragraph[start : start + chunk_chars]
                if piece.strip():
                    chunks.append(TextChunk(index=index, content=piece.strip(), page=page))
                    index += 1
            continue
        if len(buffer) + len(paragraph) + 2 > chunk_chars:
            flush()
            buffer = buffer[-overlap:] if overlap and buffer else ""
        buffer = f"{buffer}\n\n{paragraph}".strip() if buffer else paragraph

    flush()
    return chunks


async def extract(path: Path, mime: str) -> ExtractionResult:
    """Extract and chunk the text of an uploaded file.

    Args:
        path: Path to the stored file.
        mime: The file's MIME type.

    Returns:
        An :class:`ExtractionResult`. Images return an empty result -- they are
        passed to the vision model rather than extracted.

    Raises:
        ExtractionError: The file type is unsupported or unreadable.
    """
    suffix = path.suffix.lower()

    if mime.startswith("image/"):
        return ExtractionResult(text="", chunks=[], pages=0)

    if mime == "application/pdf" or suffix == ".pdf":
        text, pages = await asyncio.to_thread(_extract_pdf_sync, path)
        chunks: list[TextChunk] = []
        for number, page_text in pages:
            chunks.extend(
                chunk_text(page_text, page=number, start_index=len(chunks))
            )
        logger.info("Extracted %d chars from %d PDF page(s).", len(text), len(pages))
        return ExtractionResult(text=text, chunks=chunks, pages=len(pages))

    if suffix == ".docx" or mime in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }:
        text = await asyncio.to_thread(_extract_docx_sync, path)
        return ExtractionResult(text=text, chunks=chunk_text(text), pages=0)

    if suffix in {".txt", ".md", ".markdown"} or mime.startswith("text/"):
        text = await asyncio.to_thread(_extract_plain_sync, path)
        return ExtractionResult(text=text, chunks=chunk_text(text), pages=0)

    raise ExtractionError(f"Unsupported file type: {mime or suffix or 'unknown'}")
